"""
Automated PDF / CSV report generation for DTI / DA submission.

Implements the "Automated PDF/CSV report generation for DTI/DA submission"
core feature from the PresyoWatch scope.
"""
import csv
import io
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet


def build_csv_report(rows, date_from, date_to):
    """
    rows: list of dicts with keys:
        product, unit, city, market, price, status, is_spike, report_date, reporter
    Returns CSV bytes.
    """
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["PresyoWatch Price Report", f"{date_from} to {date_to}"])
    writer.writerow([])
    writer.writerow(
        ["Product", "Unit", "City", "Market", "Price (PHP)", "Status", "Spike?", "Date", "Reported By"]
    )
    for r in rows:
        writer.writerow(
            [
                r["product"], r["unit"], r["city"], r["market"],
                f"{r['price']:.2f}", r["status"], "YES" if r["is_spike"] else "NO",
                r["report_date"], r["reporter"],
            ]
        )
    return buf.getvalue().encode("utf-8")


def build_pdf_report(rows, date_from, date_to, city_filter="All Cities"):
    """
    rows: same structure as build_csv_report.
    Returns PDF bytes suitable for DTI/DA submission.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("<b>PresyoWatch — Market Price Report</b>", styles["Title"]))
    elements.append(Paragraph(f"Coverage: {city_filter}", styles["Normal"]))
    elements.append(Paragraph(f"Period: {date_from} to {date_to}", styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))

    table_data = [["Product", "Market", "Price (PHP)", "Status", "Spike", "Date"]]
    for r in rows:
        table_data.append(
            [
                r["product"],
                r["market"],
                f"{r['price']:.2f}",
                r["status"].capitalize(),
                "YES" if r["is_spike"] else "-",
                str(r["report_date"]),
            ]
        )

    table = Table(table_data, repeatRows=1, colWidths=[4 * cm, 4.5 * cm, 2.7 * cm, 2.5 * cm, 1.8 * cm, 2.5 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f5132")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)

    spikes = sum(1 for r in rows if r["is_spike"])
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(
        Paragraph(
            f"Total reports: {len(rows)} &nbsp;&nbsp;|&nbsp;&nbsp; Flagged possible price gouging: {spikes}",
            styles["Normal"],
        )
    )
    elements.append(
        Paragraph(
            "Generated automatically by PresyoWatch for submission to the Department of Trade "
            "and Industry (DTI) and Department of Agriculture (DA).",
            styles["Italic"],
        )
    )

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def build_xlsx_report(rows, date_from, date_to):
    """Excel version of build_csv_report. Returns .xlsx bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Price Report"
    ws.append(["PresyoWatch Price Report", f"{date_from} to {date_to}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Product", "Unit", "City", "Market", "Price (PHP)", "Status", "Spike?", "Date", "Reported By"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([
            r["product"], r["unit"], r["city"], r["market"], round(r["price"], 2),
            r["status"], "YES" if r["is_spike"] else "NO", str(r["report_date"]), r["reporter"],
        ])
    for col, width in zip("ABCDEFGHI", (24, 8, 16, 28, 12, 10, 8, 12, 24)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
